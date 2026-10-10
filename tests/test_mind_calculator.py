"""心智单元算法、导入互通、实例保存与离线扫描回归。"""
import base64
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
from PIL import Image
from openpyxl import Workbook, load_workbook

from module.api.config_service import ConfigService
from module.api.mind_calculator_service import MindCalculatorService, import_ships
from module.api.protocol import ApiError, MindCalculateParams, MindShip
from module.api.router import Router
from module.runtime.mind_calculator import MIND_COSTS, calculate, catalog, detect_rows, recognize
from module.runtime.mind_recognition import Card, ScanMerger, color_rows, estimate_scroll, level_vote, recognize_cards
from tests.test_api import fixture


def ship(name='测试舰', level=100, rarity='SSR', **kwargs):
    return MindShip(name=name, level=level, rarity=rarity, **kwargs).model_dump()


def encoded(raw):
    return base64.b64encode(raw).decode()


class CalculatorTests(unittest.TestCase):
    def test_all_rarities_and_stage_boundaries(self):
        for rarity, costs in MIND_COSTS.items():
            for level, start in [(1, 0), (99, 0), (100, 0), (101, 1), (105, 1), (106, 2),
                                 (110, 2), (111, 3), (115, 3), (116, 4), (119, 4), (120, 4), (125, 4)]:
                with self.subTest(rarity=rarity, level=level):
                    result = calculate([ship(level=level, rarity=rarity)])
                    self.assertEqual(result['mind'], sum(costs[start:]))
                    self.assertEqual(result['gold'], sum(costs[start:]) * 10)
                    self.assertEqual(result['included'], 1)

    def test_retrofit_merge_meta_and_exclusions(self):
        result = calculate([ship('拉菲', 100), ship('拉菲.改', 106), ship('皇家方舟', 110),
                            ship('皇家方舟·META', 100), ship('小企业'), ship('光辉(μ兵装)'),
                            ship('人工排除', excluded=True)])
        self.assertEqual(result['included'], 3)
        self.assertEqual(result['merged'], 1)
        self.assertEqual(result['excluded'], 3)
        row = result['ships'][1]
        self.assertEqual(row['base_rarity'], 'SR')
        self.assertEqual(row['mind'], 960)
        self.assertNotEqual(result['ships'][2]['base_name'], result['ships'][3]['name'])

    def test_unknown_retrofit_and_review_never_silently_billed(self):
        result = calculate([ship('未来舰.改'), ship('未来舰', review=True), ship('零等级', level=0)])
        self.assertEqual(result['mind'], 0)
        self.assertEqual(result['review'], 3)
        self.assertEqual(calculate([ship('未来舰.改', base_rarity='R')])['mind'], 880)

    def test_public_catalog_integrity(self):
        for name, info in catalog()['ships'].items():
            self.assertEqual(name, info['name'])
            self.assertIn(info['base_rarity'], MIND_COSTS)
        self.assertGreater(len(catalog()['ships']), 800)

    def test_rows_follow_scrolled_level_headers(self):
        ocr = Mock()
        ocr.det.return_value = [('Lv.120', [[170, 105], [218, 105], [218, 120], [170, 120]], .99),
                                ('Lv.99', [[350, 105], [398, 105], [398, 120], [350, 120]], .99),
                                ('Lv.120', [[170, 332], [218, 332], [218, 348], [170, 348]], .99)]
        self.assertEqual(detect_rows(Image.new('RGB', (1280, 720)), ocr), [105, 332])

    def test_offline_screenshot_uses_ap_models_and_marks_review(self):
        image = Image.open(Path(__file__).parent / 'fixtures/fleet_names_vanguard.png').convert('RGB')
        rows = recognize(image)
        self.assertEqual(len(rows), 21)
        self.assertEqual(rows[5]['name'], '灵敏·META')
        self.assertEqual(rows[13]['name'], '热心.改')
        self.assertEqual(rows[13]['level'], 105)
        self.assertTrue(all(row['review'] for row in rows))
        self.assertEqual([row['level'] for row in rows[:14]], [125, 100, 125, 105, 125, 63, 99, 99, 117, 122, 99, 98, 105, 105])
        self.assertEqual(rows[14]['name'], '朱诺.改')
        self.assertEqual(rows[14]['level'], 120)
        with self.assertRaises(ValueError):
            recognize(Image.new('RGB', (2560, 1440)))


class RecognitionTests(unittest.TestCase):
    def test_digit_count_rejects_lost_hundreds_and_conflicting_readings(self):
        self.assertEqual(level_vote(['Lv.25', 'Lv.125', '125'], 3)[0], 125)
        self.assertTrue(level_vote(['25', '25', '25'], 3)[1])
        self.assertEqual(level_vote(['Lv.120', 'Lv.120', '120'], 3), (120, ''))
        self.assertTrue(level_vote(['100', '101', '100'], 3)[1])
        self.assertEqual(level_vote(['noise', '126', '0'])[0], 0)

    def test_colored_header_follows_scrolling_and_unread_cards_survive(self):
        pixels = np.zeros((720, 1280, 3), dtype=np.uint8)
        pixels[103:107, 93:231] = (235, 185, 60)
        self.assertEqual(color_rows(pixels), [103])
        ocr = Mock()
        ocr.det.return_value = []
        ocr.ocr_for_single_lines.side_effect = lambda regions: [''] * len(regions)
        cards = recognize_cards(Image.fromarray(pixels), name_ocr=ocr, level_ocr=ocr)
        self.assertEqual(len(cards), 1)
        self.assertIn('未识别舰船', cards[0].ship['name'])
        self.assertTrue(cards[0].ship['review'])
        self.assertEqual(cards[0].ship['level'], 0)

    def test_third_row_is_read_and_gray_artwork_is_not_a_card_header(self):
        pixels = np.zeros((720, 1280, 3), dtype=np.uint8)
        for y in (76, 303, 530):
            pixels[y:y + 4, 93:231] = (235, 185, 60)
        # 实际错误现场中的灰色卡面条带位于列内部，不能生成一排虚假舰船。
        pixels[381:383, 444:535] = (190, 190, 190)
        self.assertEqual(color_rows(pixels), [76, 303, 530])
        ocr = Mock()
        ocr.det.return_value = []
        ocr.ocr_for_single_lines.side_effect = lambda regions: [''] * len(regions)
        cards = recognize_cards(Image.fromarray(pixels), name_ocr=ocr, level_ocr=ocr)
        self.assertEqual([card.y for card in cards], [76, 303, 530])

    def test_color_rows_survive_partial_level_detection(self):
        from module.runtime.mind_recognition import CARD_COLUMNS
        pixels = np.zeros((720, 1280, 3), dtype=np.uint8)
        # 色框圆角在不同排有像素偏差；整屏 OCR 可能只检测出其中两排的等级。
        rows = (75, 301, 528)
        for y in rows:
            for x in CARD_COLUMNS:
                pixels[y:y + 4, x:x + 138] = (235, 185, 60)
        for missing in rows:
            with self.subTest(missing=missing):
                ocr = Mock()
                ocr.det.return_value = [
                    ('Lv.120', [[170, y + 5], [218, y + 5], [218, y + 20], [170, y + 20]], .99)
                    for y in rows if y != missing]
                ocr.ocr_for_single_lines.side_effect = lambda regions: [''] * len(regions)
                cards = recognize_cards(Image.fromarray(pixels), name_ocr=ocr, level_ocr=ocr)
                self.assertEqual(sorted({card.y for card in cards}), list(rows))
                self.assertEqual(len(cards), 21)
                self.assertEqual(sum(card.y == missing for card in cards), 7)

        # 显式指定裁剪行时仍只读取调用者选择的行。
        cards = recognize_cards(Image.fromarray(pixels), name_ocr=ocr, level_ocr=ocr, row_origins=[301])
        self.assertEqual(len(cards), 7)
        self.assertEqual({card.y for card in cards}, {301})

    @staticmethod
    def scroll_frames(offsets):
        rng = np.random.default_rng(20261010)
        # 随机纹理模拟不同卡面，原图提供超过一屏的内容以核对实际位移。
        content = rng.integers(30, 230, (max(offsets) + 575, 1140, 3), dtype=np.uint8)
        frames = []
        for offset in offsets:
            frame = np.zeros((720, 1280, 3), dtype=np.uint8)
            frame[65:640, 85:1225] = content[offset:offset + 575]
            frames.append(frame)
        return frames

    @staticmethod
    def scroll_images():
        return RecognitionTests.scroll_frames([0, 227])

    def test_spatial_merge_keeps_distinct_copies_and_best_overlap(self):
        first, second = self.scroll_images()
        self.assertEqual(estimate_scroll(first, second), 227)
        self.assertEqual(estimate_scroll(second, first), -227)
        merger = ScanMerger()
        merger.add(Image.fromarray(first), [Card(93, 76, 0, ship('拉菲', 100), 6),
                                            Card(93, 303, 0, ship('拉菲', 100), 1)])
        merger.add(Image.fromarray(second), [Card(93, 76, 0, ship('拉菲', 101), 6),
                                             Card(93, 303, 0, ship('拉菲', 100), 6)])
        self.assertEqual(len(merger.ships()), 3)
        self.assertEqual(merger.ships()[1]['level'], 101)
        self.assertIn('跨屏读数不一致', merger.ships()[1]['source'])
        # 相同截图是零位移，重试截图不增加卡片。
        merger.add(Image.fromarray(second), [Card(93, 76, 0, ship('拉菲', 101), 6)])
        self.assertEqual(len(merger.ships()), 3)

    def test_unrelated_screens_are_not_guessed_as_a_scroll(self):
        first, _ = self.scroll_images()
        other = np.random.default_rng(42).integers(30, 230, (720, 1280, 3), dtype=np.uint8)
        with self.assertRaises(ValueError):
            estimate_scroll(first, other)

    def test_three_row_pages_merge_using_intermediate_scroll_evidence(self):
        frames = self.scroll_frames([0, 227, 454, 681])
        merger = ScanMerger()
        cards = [Card(93, y, 0, ship('拉菲'), 6) for y in (76, 303, 530)]
        merger.add(Image.fromarray(frames[0]), copy.deepcopy(cards))
        for frame in frames[1:]:
            merger.advance(Image.fromarray(frame))
        merger.add(Image.fromarray(frames[-1]), copy.deepcopy(cards))
        self.assertEqual(merger.offset, 681)
        self.assertEqual(len(merger.ships()), 6)

    def scan_frames(self, offsets, *, bottom_offset, rows=(76, 303, 530)):
        from module.retire.mind_scan import MindCalculatorScan
        from types import SimpleNamespace
        class ImmediateTimer:
            def __init__(self, limit, *args, **kwargs):
                self.limit = limit
            def start(self):
                return self
            def reset(self):
                return self
            def reached(self):
                return self.limit < 5
        images = self.scroll_frames(sorted(set(offsets)))
        by_offset = dict(zip(sorted(set(offsets)), images))
        frames = iter(offsets)
        scanner = MindCalculatorScan.__new__(MindCalculatorScan)
        scanner.config = SimpleNamespace(Emulator_ControlMethod='MaaTouch')
        scanner.device = Mock()
        position = [0]
        def screenshot():
            position[0] = next(frames)
            scanner.device.image = by_offset[position[0]]
        scanner.device.screenshot.side_effect = screenshot
        scanner.appear = Mock(return_value=True)
        scrollbar = SimpleNamespace(appear=lambda _: True, total=565,
                                    match_color=lambda _: np.arange(565) >= 475 if position[0] == bottom_offset
                                    else (np.arange(565) >= 2) & (np.arange(565) < 92))
        captured = []
        def recognize(*args, **kwargs):
            captured.append(position[0])
            return [Card(93, y, 0, ship('拉菲'), 6) for y in rows]
        with patch('module.retire.mind_scan.Timer', ImmediateTimer), patch('module.retire.mind_scan.DOCK_SCROLL', scrollbar), \
                patch('module.retire.mind_scan.recognize_cards', side_effect=recognize):
            result = scanner._scan_pages(Mock(), Mock())
        return scanner, captured, result

    def test_scan_moves_three_rows_and_waits_for_motion_to_settle(self):
        scanner, captured, result = self.scan_frames([0, 0, 100, 227, 227, 454, 454, 681, 681], bottom_offset=681)
        self.assertEqual(captured, [0, 681])
        self.assertEqual(len(result), 6)
        self.assertEqual(scanner.device.drag.call_count, 3)
        self.assertTrue(all(call.kwargs['hold_duration'] == .4 for call in scanner.device.drag.call_args_list))
        scanner.device.swipe.assert_not_called()

    def test_color_band_offsets_do_not_accumulate_into_three_row_distance(self):
        scanner, captured, result = self.scan_frames(
            [0, 0, 227, 227, 454, 454, 681, 681, 908, 908, 1135, 1135, 1362, 1362],
            bottom_offset=1362, rows=(75, 301, 528))
        self.assertEqual(captured, [0, 681, 1362])
        self.assertEqual(len(result), 9)
        self.assertEqual(scanner.device.drag.call_count, 6)
        self.assertEqual([call.args[0][1] - call.args[1][1] for call in scanner.device.drag.call_args_list],
                         [227] * 6)

    def test_scan_corrects_overshoot_before_reading_and_allows_partial_last_page(self):
        scanner, captured, result = self.scan_frames(
            [0, 0, 240, 240, 467, 467, 690, 690, 681, 681, 908, 908], bottom_offset=908)
        self.assertEqual(captured, [0, 681, 908])
        self.assertEqual(len(result), 7)
        correction = scanner.device.drag.call_args_list[3]
        self.assertGreater(correction.args[1][1], correction.args[0][1])

    def test_adb_fallback_is_slow_and_never_falls_back_to_clicking_a_ship(self):
        from module.retire.mind_scan import MindCalculatorScan
        from module.exception import MindCalculatorScanError
        from types import SimpleNamespace
        scanner = MindCalculatorScan.__new__(MindCalculatorScan)
        scanner.config = SimpleNamespace(Emulator_ControlMethod='ADB')
        scanner.device = Mock()
        scanner._drag_rows(227)
        scanner.device.drag.assert_not_called()
        self.assertEqual(scanner.device.swipe.call_args.kwargs['duration'], .6)
        with self.assertRaises(MindCalculatorScanError):
            scanner._drag_rows(4)
        self.assertEqual(scanner.device.swipe.call_count, 1)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.configs = ConfigService(fixture(self.temp.name))
        self.service = MindCalculatorService(self.configs)

    def test_save_conflict_and_instances_are_separate(self):
        report = self.service.report('testpilot')
        saved = self.service.save('testpilot', report['revision'], [ship()])
        self.assertEqual(saved['mind'], 2200)
        self.assertEqual(self.service.report('testpilot')['ships'][0]['name'], '测试舰')
        with self.assertRaises(ApiError) as caught:
            self.service.save('testpilot', report['revision'], [])
        self.assertEqual(caught.exception.code, 'CONFLICT')
        self.configs.create('other')
        self.assertEqual(self.service.report('other')['ships'], [])

    def test_all_exports_import_back_without_losing_review_or_exclusion(self):
        rows = [ship('拉菲.改', 105, source='截图.png'), ship('待核对舰', review=True), ship('排除舰', excluded=True)]
        self.service.save('testpilot', self.service.report('testpilot')['revision'], rows)
        for format in ['json', 'csv', 'xlsx']:
            file = self.service.export('testpilot', format)
            restored = import_ships(file['filename'], file['content'])
            self.assertEqual(restored[0]['level'], 105)
            self.assertEqual(restored[0]['base_rarity'], 'SR')
            self.assertTrue(restored[1]['review'])
            self.assertTrue(restored[2]['excluded'])

    def test_import_original_excel_header_at_row_two(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '舰船明细'
        sheet.append(['搜索船名'])
        sheet.append(['序号', '船名', '舰种', '等级', '稀有度', '状态'])
        sheet.append([1, '独角兽.改', '航母', 115, '精锐', ''])
        output = io.BytesIO()
        book.save(output)
        rows = import_ships('原工具.xlsx', encoded(output.getvalue()))
        self.assertEqual(calculate(rows)['mind'], 600)
        sheet['F3'] = '存疑'
        sheet['M2'], sheet['M3'] = '排除船名', '小企业'
        output = io.BytesIO()
        book.save(output)
        rows = import_ships('原工具.xlsx', encoded(output.getvalue()))
        self.assertTrue(rows[0]['review'])
        self.assertTrue(rows[1]['excluded'])
        self.assertEqual(calculate(rows)['mind'], 0)

    def test_reject_invalid_import_and_formula_cells_are_text(self):
        with self.assertRaises(ApiError):
            import_ships('bad.json', encoded(json.dumps([{'name': '错误', 'level': 126}]).encode()))
        with self.assertRaises(ApiError):
            import_ships('bad.csv', encoded('船名,等级\n测试,100.1'.encode()))
        self.service.save('testpilot', self.service.report('testpilot')['revision'], [ship('=1+1')])
        file = self.service.export('testpilot', 'xlsx')
        book = load_workbook(io.BytesIO(base64.b64decode(file['content'])))
        self.assertEqual(book['舰船明细']['A2'].data_type, 's')
        file = self.service.export('testpilot', 'csv')
        self.assertIn("'=1+1", base64.b64decode(file['content']).decode('utf-8-sig'))
        self.assertEqual(import_ships(file['filename'], file['content'])[0]['name'], '=1+1')

    def test_rpc_contract_and_demo_read_only(self):
        router = Router(self.configs, Mock())
        self.assertEqual(router.dispatch('mind.calculate', {'instance': 'testpilot', 'ships': [ship()]})['mind'], 2200)
        with patch.dict('os.environ', {'DEMO': '1'}), self.assertRaises(ApiError):
            router.dispatch('mind.save', {'instance': 'testpilot', 'ships': [], 'revision': self.service.report('testpilot')['revision']})
        with self.assertRaises(ValueError):
            MindCalculateParams(instance='testpilot', ships=[dict(name='测试', level=True)])

    def test_completed_scan_preserves_concurrent_data_and_same_list_save(self):
        from module.config.config import AzurLaneConfig
        from module.exception import RequestHumanTakeover
        from module.retire.mind_scan import MindCalculatorScan
        for changed in (False, True):
            with self.subTest(changed=changed):
                self.service.save('testpilot', self.service.report('testpilot')['revision'], [])
                path = str(self.configs.path('testpilot'))
                with patch('module.config.utils.filepath_config', return_value=path), \
                        patch('module.config.config.filepath_config', return_value=path), \
                        patch('module.config.config_updater.filepath_config', return_value=path), \
                        patch('module.retire.mind_scan.DOCK_SCROLL') as scrollbar, \
                        patch('module.retire.mind_scan.server.server', 'cn'):
                    config = AzurLaneConfig.__new__(AzurLaneConfig)
                    config.config_name = 'testpilot'
                    config.data = config.read_file('testpilot')
                    config._loaded_data = copy.deepcopy(config.data)
                    config.modified = {}
                    scanner = MindCalculatorScan.__new__(MindCalculatorScan)
                    scanner.config = config
                    scanner.ui_ensure, scanner.dock_reset = Mock(), Mock()
                    scanner.appear = Mock(return_value=False)
                    scrollbar.appear.return_value = scrollbar.at_top.return_value = True
                    def collect(*args):
                        self.service.save('testpilot', self.service.report('testpilot')['revision'],
                                          [ship('用户的新数据')] if changed else [])
                        return [ship('拉菲')]
                    scanner._scan_pages = Mock(side_effect=collect)
                    if changed:
                        with self.assertRaises(RequestHumanTakeover):
                            scanner.run()
                    else:
                        scanner.run()
                    rows = self.service.report('testpilot')['ships']
                    self.assertEqual(rows[0]['name'], '用户的新数据' if changed else '拉菲')


if __name__ == '__main__':
    unittest.main()
